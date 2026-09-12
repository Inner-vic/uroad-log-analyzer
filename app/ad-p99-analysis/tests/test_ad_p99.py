import contextlib
import io
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import ad_p99


ORDER = "slow_driver/parking/HNOA/PGear/fast_driver/CNOA/switch_parking_driving/user_RGear"


def sample_book():
    """Synthetic workbook based on the user screenshots, not an upstream export."""
    book = Workbook()
    vla = book.active
    vla.title = "VLA"
    parking = book.create_sheet("VLA Parking")
    for sheet, metric, first, values in (
        (vla, "header_delay", 52, "281.16/-/-/-/194.0/180.0/-/-"),
        (parking, "vla_parking_header_delay", 72, "609.6/377.0/-/-/-/-/-/-"),
    ):
        sheet.merge_cells(start_row=first, start_column=2, end_row=first + 4, end_column=2)
        sheet.cell(first, 2, f"{metric}(ms)\n({ORDER})")
        for offset, statistic in enumerate(("mean", "max", "min", "P90", "P99")):
            sheet.cell(first + offset, 3, statistic)
            sheet.cell(first + offset, 4, values if statistic == "P99" else "999/999/999/999/999/999/999/999")
        # A neighboring P99 must never be mistaken for the header metric.
        sheet.cell(first + 5, 2, f"{metric.replace('header', 'pipeline')}(ms)\n({ORDER})")
        sheet.cell(first + 5, 3, "P99")
        sheet.cell(first + 5, 4, "1/2/3/4/5/6/7/8")
    return book


def as_bytes(book):
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def parse(book):
    return ad_p99.parse_workbook(as_bytes(book))


class ParserTests(unittest.TestCase):
    def test_compact_ignores_invisible_sheet_title_characters(self):
        self.assertEqual(ad_p99.compact("VLA\u200b"), "vla")
        self.assertEqual(ad_p99.compact("VLA\ufeff Parking"), "vlaparking")

    def test_screenshot_values_and_evidence(self):
        data = as_bytes(sample_book())
        result = ad_p99.parse_workbook(data, metadata={"release": "synthetic-example"})
        self.assertEqual(result["status"], "ok")
        self.assertEqual({key: value["p99_ms"] for key, value in result["metrics"].items()}, {
            "slow_driver": 281.16, "fast_driver": 194.0, "CNOA": 180.0, "parking": 377.0,
        })
        self.assertEqual(result["metrics"]["parking"]["evidence"]["value_cell"], "D76")
        self.assertEqual(result["metrics"]["slow_driver"]["evidence"]["value_cell"], "D56")
        self.assertEqual(result["metrics"]["parking"]["evidence"]["scenario_index"], 1)
        self.assertEqual(result["metadata"]["release"], "synthetic-example")
        self.assertEqual(len(result["source"]["sha256"]), 64)

    def test_reordered_scenarios_moved_rows_and_case(self):
        book = sample_book()
        sheet = book["VLA"]
        sheet.unmerge_cells("B52:B56")
        sheet.title = " vLa "
        sheet["B52"] = "header_delay（ms）\n(CNOA/fast_driver/slow_driver)"
        sheet["D56"] = "180/194/281.16"
        sheet.move_range("B52:D57", rows=-40, cols=2)
        result = parse(book)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["metrics"]["slow_driver"]["p99_ms"], 281.16)
        self.assertEqual(result["metrics"]["slow_driver"]["evidence"]["value_cell"], "F16")

    def test_missing_value_is_null_and_zero_is_valid(self):
        book = sample_book()
        book["VLA"]["D56"] = "0/-/-/-/-/180/-/-"
        result = parse(book)
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["metrics"]["slow_driver"]["p99_ms"], 0)
        self.assertIsNone(result["metrics"]["fast_driver"]["p99_ms"])
        self.assertEqual(result["metrics"]["parking"]["p99_ms"], 377)

    def test_missing_sheet_preserves_other_metrics(self):
        book = sample_book()
        del book["VLA Parking"]
        result = parse(book)
        self.assertEqual(result["status"], "partial")
        self.assertIsNone(result["metrics"]["parking"]["p99_ms"])
        self.assertEqual(result["metrics"]["slow_driver"]["p99_ms"], 281.16)

    def test_ambiguous_sheet_is_not_selected(self):
        book = sample_book()
        book.create_sheet("VLA  Parking")
        result = parse(book)
        self.assertIsNone(result["metrics"]["parking"]["p99_ms"])

    def test_does_not_fall_through_to_pipeline_p99(self):
        for merged in (True, False):
            with self.subTest(merged=merged):
                book = sample_book()
                if not merged:
                    book["VLA"].unmerge_cells("B52:B56")
                book["VLA"]["C56"] = "other"
                result = parse(book)
                self.assertIsNone(result["metrics"]["slow_driver"]["p99_ms"])

    def test_duplicate_metric_or_p99_is_ambiguous(self):
        for target, value in (("B60", f"header_delay(ms)({ORDER})"), ("C54", "P99")):
            with self.subTest(target=target):
                book = sample_book()
                book["VLA"][target] = value
                result = parse(book)
                self.assertIsNone(result["metrics"]["slow_driver"]["p99_ms"])

    def test_mismatched_count_and_duplicate_labels(self):
        for target, value in (("D56", "281.16/194/180"),
                              ("B52", "header_delay(ms)(slow_driver/slow_driver)")):
            with self.subTest(target=target):
                book = sample_book()
                book["VLA"][target] = value
                result = parse(book)
                self.assertIsNone(result["metrics"]["slow_driver"]["p99_ms"])

    def test_reject_bad_units_nonfinite_and_negative(self):
        for target, value in (("B52", f"header_delay(s)({ORDER})"),
                              ("D56", "nan/-/-/-/inf/-3/-/-")):
            with self.subTest(target=target):
                book = sample_book()
                book["VLA"][target] = value
                result = parse(book)
                self.assertTrue(all(result["metrics"][key]["p99_ms"] is None
                                    for key in ("slow_driver", "fast_driver", "CNOA")))

    def test_formula_without_cache_is_not_evaluated(self):
        book = sample_book()
        book["VLA"]["D56"] = '=CONCAT("281.16/-/-/-/194/180/-/-")'
        result = parse(book)
        self.assertIsNone(result["metrics"]["slow_driver"]["p99_ms"])
        self.assertIn("缓存", result["issues"][0]["message"])

    def test_merged_statistic_columns(self):
        book = sample_book()
        sheet = book["VLA"]
        sheet["E56"] = sheet["D56"].value
        sheet.merge_cells("C56:D56")
        result = parse(book)
        self.assertEqual(result["metrics"]["slow_driver"]["p99_ms"], 281.16)
        self.assertEqual(result["metrics"]["slow_driver"]["evidence"]["value_cell"], "E56")

    def test_missing_scenario_cannot_be_replaced_by_hnoa(self):
        book = sample_book()
        book["VLA"]["B52"] = book["VLA"]["B52"].value.replace("CNOA", "another_mode")
        result = parse(book)
        self.assertIsNone(result["metrics"]["CNOA"]["p99_ms"])
        self.assertEqual(result["metrics"]["fast_driver"]["p99_ms"], 194)

    def test_invalid_file_and_no_target_sheets(self):
        with self.assertRaises(ad_p99.InputError):
            ad_p99.parse_workbook(b"<html>login</html>")
        self.assertEqual(parse(Workbook())["status"], "failed")


class IntakeTests(unittest.TestCase):
    def test_nested_card_url_and_callback(self):
        card = {"elements": [
            {"tag": "button", "text": {"content": "Log下载"}, "url": "https://example.test/log"},
            {"tag": "button", "text": {"content": "查看结果"},
             "multi_url": {"url": "https://example.test/report", "pc_url": "https://example.test/report"}},
            {"tag": "button", "text": "查看结果", "value": {"secret": "not-output"}},
        ]}
        result = ad_p99.inspect_card({"body": {"content": json.dumps(card)}})
        self.assertEqual(result["result_buttons"], [
            {"urls": ["https://example.test/report"], "has_callback": False},
            {"urls": [], "has_callback": True},
        ])
        self.assertNotIn("not-output", json.dumps(result))

    def test_v2_card_and_missing_button(self):
        result = ad_p99.inspect_card({"tag": "button", "text": "查看结果", "behaviors": [
            {"type": "open_url", "default_url": "https://example.test/result.xlsx"},
            {"type": "callback", "value": "ignored"},
        ]})
        self.assertEqual(result["result_buttons"][0]["urls"], ["https://example.test/result.xlsx"])
        self.assertTrue(result["result_buttons"][0]["has_callback"])
        self.assertEqual(ad_p99.inspect_card({})["result_buttons"], [])

    def test_url_boundaries(self):
        for url in ("http://example.test/x", "file:///tmp/x", "https://u:p@example.test/x",
                    "https://example.test:bad/x"):
            with self.subTest(url=url), self.assertRaises(ad_p99.InputError):
                ad_p99.checked_url(url)

    def test_card_exit_code_and_output_refuses_overwrite(self):
        with patch("ad_p99.read_local", return_value=b"{}"), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(ad_p99.main(["--card-json", "synthetic.json"]), 2)
        with patch("ad_p99.read_local") as read, contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(ad_p99.main(["--card-json", __file__, "--output", __file__]), 1)
            read.assert_not_called()

    def test_file_inputs_removed_from_cloud_entry(self):
        for flag in ("--excel", "--result-url"):
            with self.subTest(flag=flag), contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                ad_p99.main([flag, "not-a-real-input"])
        self.assertFalse(hasattr(ad_p99, "download_result"))


class ComparisonTests(unittest.TestCase):
    def test_copied_request_accepts_one_or_two_matching_vins_and_sorts_windows(self):
        request = """刷包后 VIN：DEM0VEH1CLE000001，20260402150000/20260402160000
        Baseline VIN: dem0veh1cle000001，2026-04-01 12:00:00 - 2026-04-01 13:00:00"""
        result = ad_p99.parse_comparison_request(request)
        self.assertEqual(result["status"], "ready")
        self.assertEqual(result["vin"], "DEM0VEH1CLE000001")
        self.assertEqual(result["baseline_metadata"]["start"], "2026-04-01T12:00:00")
        self.assertEqual(result["replacement_metadata"]["start"], "2026-04-02T15:00:00")
        one_vin = ad_p99.parse_comparison_request(
            "VIN：DEM0VEH1CLE000001；20260402150000/20260402160000；20260401120000/20260401130000"
        )
        self.assertEqual(one_vin["baseline_metadata"]["vin"], one_vin["replacement_metadata"]["vin"])

    def test_copied_request_rejects_cross_vin_overlapping_or_invalid_windows(self):
        with self.assertRaisesRegex(ad_p99.InputError, "多个 VIN"):
            ad_p99.parse_comparison_request(
                "Baseline VIN: DEM0VEH1CLE000001 20260401120000/20260401130000; "
                "替换后 VIN: DEM0VEH1CLE000002 20260402120000/20260402130000"
            )
        with self.assertRaisesRegex(ad_p99.InputError, "重叠"):
            ad_p99.parse_comparison_request(
                "VIN: DEM0VEH1CLE000001 20260401120000/20260401140000; 20260401130000/20260401150000"
            )
        with self.assertRaisesRegex(ad_p99.InputError, "起点必须早于终点"):
            ad_p99.parse_comparison_request(
                "VIN: DEM0VEH1CLE000001 20260401140000/20260401130000; 20260402120000/20260402130000"
            )

    def test_comparison_request_cli_only_normalizes_context(self):
        copied = b"VIN: DEM0VEH1CLE000001 20260402120000/20260402130000; 20260401120000/20260401130000"
        stdout = io.StringIO()
        with patch("ad_p99.read_local", return_value=copied), contextlib.redirect_stdout(stdout):
            self.assertEqual(ad_p99.main(["--comparison-request", "synthetic.txt"]), 0)
        payload = json.loads(stdout.getvalue())
        self.assertEqual(payload["kind"], "ad_p99_comparison_request")
        self.assertEqual(payload["replacement_metadata"]["start"], "2026-04-02T12:00:00")

    def test_compares_four_p99_metrics_and_exposes_table_rows(self):
        baseline = parse(sample_book())
        replacement_book = sample_book()
        replacement_book["VLA"]["D56"] = "280/-/-/-/195/178/-/-"
        replacement_book["VLA Parking"]["D76"] = "600/400/-/-/-/-/-/-"
        replacement = parse(replacement_book)
        context = {"vin": "VIN-001", "start": "2026-09-01T00:00:00+08:00", "end": "2026-09-01T01:00:00+08:00"}
        result = ad_p99.compare_results(
            baseline, replacement,
            baseline_metadata=context,
            replacement_metadata={**context, "start": "2026-09-02T00:00:00+08:00", "end": "2026-09-02T01:00:00+08:00"},
        )
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["comparisons"]["slow_driver"]["delta_ms"], -1.16)
        self.assertEqual(result["comparisons"]["fast_driver"]["delta_ms"], 1.0)
        self.assertEqual(result["comparisons"]["CNOA"]["delta_ms"], -2.0)
        self.assertEqual(result["comparisons"]["parking"]["delta_ms"], 23.0)
        self.assertEqual(result["table"]["columns"], ["场景", "指标", "Baseline", "替换小包", "Δ（替换 − Baseline）"])
        self.assertEqual(result["table"]["rows"][0][0], "slow_driver（低速人驾）")

    def test_requires_matching_vin_and_complete_time_ranges(self):
        result = parse(sample_book())
        with self.assertRaisesRegex(ad_p99.InputError, "vin 不一致"):
            ad_p99.compare_results(
                result, result,
                baseline_metadata={"vin": "VIN-001", "start": "a", "end": "b"},
                replacement_metadata={"vin": "VIN-002", "start": "c", "end": "d"},
            )
        with self.assertRaisesRegex(ad_p99.InputError, "缺少非空 end"):
            ad_p99.compare_results(
                result, result,
                baseline_metadata={"vin": "VIN-001", "start": "a", "end": "b"},
                replacement_metadata={"vin": "VIN-001", "start": "c"},
            )


if __name__ == "__main__":
    unittest.main()
