"""Read Lark Sheets into memory through the installed, authenticated CLI."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import urllib.parse
from pathlib import Path

from ad_p99 import InputError, MAX_SHEET_CELLS, checked_url, compact, parse_loaded_book

READ_COMMANDS = {"+workbook-info", "+sheet-info", "+csv-get", "+revision-get"}


def cli_executable(explicit=None):
    candidate = explicit or shutil.which("lark-cli")
    if not candidate:
        raise InputError("未找到 lark-cli；请由部署人员安装并配置飞书应用。")
    path = Path(candidate)
    if os.name == "nt" and path.suffix.lower() != ".exe":
        # Resolve the npm shim to its native binary, without invoking a shell.
        path = path.parent / "node_modules" / "@larksuite" / "cli" / "bin" / "lark-cli.exe"
    if not path.is_file():
        raise InputError("未找到 lark-cli 原生可执行文件；请用 --lark-cli 指定路径。")
    return str(path)


class LarkReader:
    def __init__(self, url, identity="bot", executable=None):
        checked_url(url)
        if not re.fullmatch(r"/(sheets|spreadsheets|wiki)/[A-Za-z0-9_-]+/?", urllib.parse.urlsplit(url).path):
            raise InputError("需要飞书电子表格链接（/sheets/ 或指向表格的 /wiki/），普通 Docx 暂不支持。")
        if identity not in ("bot", "user"):
            raise InputError("飞书身份必须为 bot 或 user。")
        self.url = url
        self.identity = identity
        self.executable = cli_executable(executable)

    def call(self, command, *flags):
        if command not in READ_COMMANDS:
            raise InputError("在线数据源仅允许只读命令。")
        env = dict(os.environ, LARKSUITE_CLI_NO_UPDATE_NOTIFIER="1", LARKSUITE_CLI_NO_SKILLS_NOTIFIER="1")
        try:
            completed = subprocess.run(
                [self.executable, "sheets", command, "--url", self.url,
                 "--as", self.identity, "--format", "json", *flags],
                capture_output=True, text=True, encoding="utf-8", timeout=60, env=env,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
            )
        except (OSError, subprocess.TimeoutExpired):
            raise InputError("飞书读取进程启动失败或超过 60 秒；请检查 CLI 配置和网络。") from None
        try:
            payload = json.loads(completed.stdout if completed.returncode == 0 else completed.stderr)
        except (ValueError, TypeError):
            raise InputError("lark-cli 未返回有效 JSON；请检查版本及配置。") from None
        if not isinstance(payload, dict):
            raise InputError("lark-cli 返回结构不受支持。")
        if completed.returncode or payload.get("ok") is not True:
            error = payload.get("error", {})
            if isinstance(error, dict) and error.get("subtype") == "not_configured":
                raise InputError("lark-cli 尚未配置飞书应用；需要部署人员完成应用配置及只读访问授权。")
            if isinstance(error, dict) and error.get("type") == "authorization":
                raise InputError(f"飞书 {self.identity} 身份无权读取；请检查只读 scope 和文档访问权限。")
            raise InputError(f"飞书只读命令 {command} 失败；请检查文档类型、权限、链接和 CLI 状态。")
        data = payload.get("data")
        if not isinstance(data, dict):
            raise InputError("lark-cli data 不是对象；需要按实际版本适配。")
        return data


def require_complete(data):
    if data.get("has_more") or data.get("truncated") or data.get("complete") is False:
        raise InputError("飞书响应被截断，不能以部分数据作为完整结果。")


def read_online(url, *, identity="bot", executable=None, metadata=None, reader=None):
    """No export, local input file, CSV file, or XLSX roundtrip is involved."""
    try:
        from openpyxl import Workbook
        from openpyxl.utils.cell import get_column_letter, column_index_from_string, range_boundaries
    except ImportError:
        raise InputError("缺少 openpyxl；请由部署人员按 requirements.txt 准备依赖。") from None
    reader = reader or LarkReader(url, identity, executable)
    initial_revision = reader.call("+revision-get").get("revision")
    if not isinstance(initial_revision, (str, int)) or isinstance(initial_revision, bool) or initial_revision == "":
        raise InputError("未取得表格 revision，无法核对读取期间的版本一致性。")
    info = reader.call("+workbook-info")
    require_complete(info)
    if info.get("warning_message"):
        raise InputError("工作簿清单有警告；请核对 sheet 是否被删除或更名后再读取。")
    sheets = info.get("sheets")
    if sheets is None and isinstance(info.get("workbook"), dict):
        sheets = info["workbook"].get("sheets")
    if not isinstance(sheets, list) or not all(isinstance(item, dict) for item in sheets):
        raise InputError("工作簿清单缺少 sheets 数组。")
    targets = [item for item in sheets if compact(item.get("title")) in ("vla", "vlaparking")]
    names = [compact(item.get("title")) for item in targets]
    if len(names) != len(set(names)):
        raise InputError("存在多个同名目标 sheet，无法唯一确定结果。")
    book = Workbook()
    book.remove(book.active)
    snapshot = hashlib.sha256()
    ids, warnings = {}, []
    try:
        for item in targets:
            title, sid = item.get("title"), item.get("sheet_id")
            row_count, col_count = item.get("row_count"), item.get("column_count")
            if not isinstance(sid, str) or not sid:
                raise InputError("目标工作表缺少 sheet_id。")
            if any(type(value) is not int or value < 1 for value in (row_count, col_count)):
                raise InputError("目标工作表缺少有效行列边界。")
            if row_count * col_count > MAX_SHEET_CELLS:
                raise InputError("在线目标 sheet 超过 500000 格扫描上限；需按实际布局缩小读取策略。")
            sheet = book.create_sheet(title)
            ids[title] = sid
            layout = reader.call("+sheet-info", "--sheet-id", sid, "--include", "merges")
            require_complete(layout)
            if layout.get("warning_message"):
                warnings.append(str(layout["warning_message"]))

            def read_window(first, last):
                requested = f"A{first}:{get_column_letter(col_count)}{last}"
                page = reader.call("+csv-get", "--sheet-id", sid, "--range", requested,
                                   "--include-row-prefix=false", "--skip-hidden=false", "--max-chars", "500000")
                # Never consume a truncated page; re-read disjoint smaller windows.
                if page.get("has_more") or page.get("truncated") or page.get("complete") is False:
                    if first == last:
                        raise InputError("单行读取仍被截断，无法完整获取 P99 数据。")
                    middle = (first + last) // 2
                    read_window(first, middle)
                    read_window(middle + 1, last)
                    return
                if page.get("warning_message"):
                    warnings.append(str(page["warning_message"]))
                rows, cols = page.get("row_indices"), page.get("col_indices")
                if rows != list(range(first, last + 1)) or cols != [get_column_letter(n) for n in range(1, col_count + 1)]:
                    raise InputError("在线读取返回行列与请求范围不一致，拒绝以不完整数据解析。")
                if not isinstance(page.get("annotated_csv"), str):
                    raise InputError("在线读取缺少 annotated_csv 数据。")
                try:
                    values = list(csv.reader(io.StringIO(page["annotated_csv"]), strict=True))
                except csv.Error:
                    raise InputError("在线单元格数据无法按 CSV 正确解析。") from None
                if len(values) != len(rows) or any(len(row) != len(cols) for row in values):
                    raise InputError("在线单元格矩阵与服务端坐标不一致。")
                for row_number, row in zip(rows, values):
                    for column, value in zip(cols, row):
                        if value != "":
                            cell = sheet.cell(row_number, column_index_from_string(column))
                            cell.value = value
                            cell.data_type = "s"  # Literal server values; never execute formula text.
                snapshot.update(json.dumps([sid, requested, values], ensure_ascii=False).encode("utf-8"))

            for first in range(1, row_count + 1, 200):
                read_window(first, min(first + 199, row_count))
            merges = layout.get("merged_cells", layout.get("merges"))
            if not isinstance(merges, list):
                raise InputError("在线 sheet 缺少合并区域列表，无法确认 header 指标边界。")
            for merged in merges:
                ref = merged if isinstance(merged, str) else (
                    merged.get("range") or merged.get("a1_range") or merged.get("range_ref")
                ) if isinstance(merged, dict) else None
                if not isinstance(ref, str) or not re.fullmatch(r"[A-Z]+[1-9]\d*:[A-Z]+[1-9]\d*", ref):
                    raise InputError("合并区域格式不受支持，需要按实际 CLI 返回结构适配。")
                left, top, right, bottom = range_boundaries(ref)
                if not (1 <= left <= right <= col_count and 1 <= top <= bottom <= row_count):
                    raise InputError("合并区域超出读取范围。")
                sheet.merge_cells(ref)
            snapshot.update(json.dumps([title, merges], ensure_ascii=False).encode("utf-8"))
        final_revision = reader.call("+revision-get").get("revision")
        if str(final_revision) != str(initial_revision):
            raise InputError("读取期间飞书表格版本发生变化；请等上游结果稳定后重新读取。")
        parsed_url = urllib.parse.urlsplit(url)
        clean_url = urllib.parse.urlunsplit((parsed_url.scheme, parsed_url.netloc, parsed_url.path, "", ""))
        result = parse_loaded_book(book, source={
            "kind": "lark_sheet", "url": clean_url, "revision": initial_revision,
            "snapshot_sha256": snapshot.hexdigest(), "identity": identity,
        }, metadata=metadata)
        # CLI warnings may contain source data; retain counts, not arbitrary server text.
        result["read_warning_count"] = len(warnings)
        for metric in result["metrics"].values():
            if metric["evidence"]:
                original = metric["evidence"]
                metric["evidence"] = {key: original[key] for key in (
                    "sheet", "header_cell", "statistic_cell", "value_cell", "scenario_index"
                ) if key in original}
                metric["evidence"]["sheet_id"] = ids[original["sheet"]]
        return result
    finally:
        book.close()
