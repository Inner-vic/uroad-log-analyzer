"""Independent AD P99 result intake and extraction; never runs upstream analysis."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import sys
import unicodedata
import urllib.parse
import zipfile
from decimal import Decimal
from datetime import datetime, timezone
from pathlib import Path

MAX_BYTES = 64 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_SHEET_CELLS = 500_000
SCENARIOS = (
    ("slow_driver", "低速人驾", "vla", "header_delay"),
    ("fast_driver", "高速人驾", "vla", "header_delay"),
    ("CNOA", "城区智驾", "vla", "header_delay"),
    ("parking", "泊车", "vlaparking", "vla_parking_header_delay"),
)
COMPARISON_CONTEXT_FIELDS = ("vin", "start", "end")
TIME_TOKEN = (
    r"(?:\d{14}|\d{12}|\d{4}[-/.]\d{1,2}[-/.]\d{1,2}[ T]"
    r"\d{1,2}:\d{2}(?::\d{2})?(?:Z|[+-]\d{2}:?\d{2})?)"
)
TIME_RANGE_PATTERN = re.compile(
    rf"(?P<start>{TIME_TOKEN})\s*(?:/|,|，|~|～|至|到|—|–|-)+\s*(?P<end>{TIME_TOKEN})"
)
LABELED_VIN_PATTERN = re.compile(r"(?i)\bvin(?:码|号)?\s*[:：=]?\s*([A-Za-z0-9_-]{5,})")
UNLABELED_VIN_PATTERN = re.compile(r"(?<![A-Za-z0-9])[A-HJ-NPR-Za-hj-npr-z0-9]{17}(?![A-Za-z0-9])")


class InputError(ValueError):
    pass


def compact(value):
    """Canonicalize labels received from Sheets without changing visible text semantics."""
    normalized = unicodedata.normalize("NFKC", str(value if value is not None else ""))
    return "".join(
        char for char in normalized
        if not char.isspace() and unicodedata.category(char) not in {"Cf", "Cc"}
    ).casefold()


def read_local(path):
    with Path(path).open("rb") as stream:
        data = stream.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise InputError("输入超过 64 MiB 上限。")
    return data


def checked_url(url):
    try:
        parsed = urllib.parse.urlsplit(url)
        port = parsed.port
    except ValueError as exc:
        raise InputError("结果 URL 格式无效。") from exc
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise InputError("结果地址必须是无内嵌用户名/密码的 HTTPS URL。")
    if any(ord(char) < 33 for char in url):
        raise InputError("结果 URL 含空白或控制字符。")
    return parsed.hostname.lower(), port or 443


def load_excel(data):
    """Synthetic development fixtures only; not exposed by the cloud Skill CLI."""
    if len(data) > MAX_BYTES:
        raise InputError("输入超过 64 MiB 上限。")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise InputError("输入不是 Excel OOXML 工作簿，请提供 .xlsx 文件。")
            if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise InputError("工作簿解压大小超过 256 MiB 上限。")
    except zipfile.BadZipFile:
        raise InputError("输入不是有效 .xlsx；网页、旧版 .xls 和加密文件暂不支持。") from None
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise InputError("缺少 openpyxl；请由部署人员按 requirements.txt 准备依赖。") from None
    try:
        # Cached formula values are deliberately used; the source is never saved.
        return load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    except Exception as exc:
        raise InputError(f"无法读取 Excel 工作簿（{type(exc).__name__}）。") from None


def header_labels(value, metric):
    text = compact(value)
    match = re.fullmatch(re.escape(metric) + r"\(ms\)\(([^()]*)\)", text)
    if match is None:
        raise InputError("header 格式不受支持；需包含指标名、(ms) 和括号内的场景列表。")
    labels = match.group(1).split("/")
    if not all(labels) or len(set(labels)) != len(labels):
        raise InputError("header 的场景名称为空或重复。")
    return labels


def find_block(sheet, metric):
    if sheet.max_row * sheet.max_column > MAX_SHEET_CELLS:
        raise InputError("目标 sheet 的扫描范围超过 500000 个单元格。")
    matches = [cell for row in sheet.iter_rows() for cell in row
               if isinstance(cell.value, str)
               and re.match(re.escape(metric) + r"\(", compact(cell.value))]
    if not matches:
        raise InputError(f"未找到 {metric} 指标。")
    if len(matches) != 1:
        raise InputError(f"存在多个 {metric} 指标块，无法唯一确定结果。")
    anchor = matches[0]
    labels = header_labels(anchor.value, metric)
    end_row, end_col = sheet.max_row, anchor.column
    for merged in sheet.merged_cells.ranges:
        if anchor.coordinate in merged:
            end_row, end_col = merged.max_row, merged.max_col
            break
    else:
        # Unmerged exports leave the metric cell blank on following statistic rows.
        for row in range(anchor.row + 1, sheet.max_row + 1):
            if sheet.cell(row, anchor.column).value is not None:
                end_row = row - 1
                break
    candidates = []
    for row in sheet.iter_rows(min_row=anchor.row, max_row=end_row,
                               min_col=end_col + 1):
        candidates.extend(cell for cell in row if compact(cell.value) == "p99")
    if len(candidates) != 1:
        raise InputError("header 指标块中没有唯一的 P99 行。")
    stat = candidates[0]
    value_col = stat.column + 1
    for merged in sheet.merged_cells.ranges:
        if stat.coordinate in merged:
            value_col = merged.max_col + 1
            break
    value_cell = sheet.cell(stat.row, value_col)
    if value_cell.value is None:
        raise InputError("P99 值为空或缓存结果缺失；请上游确认分析已完成并保存结果表。")
    values = compact(value_cell.value).split("/")
    if len(values) != len(labels):
        raise InputError(f"场景数量 {len(labels)} 与 P99 数值数量 {len(values)} 不一致。")
    evidence = {
        "sheet": sheet.title,
        "header_cell": anchor.coordinate,
        "statistic_cell": stat.coordinate,
        "value_cell": value_cell.coordinate,
        "raw_header": str(anchor.value),
        "raw_p99": str(value_cell.value),
        "scenario_order": labels,
    }
    return labels, values, evidence


def parse_workbook(data, *, source_name="workbook.xlsx", metadata=None):
    """Stable adapter boundary: authorized upstream clients supply Excel bytes."""
    book = load_excel(data)
    return parse_loaded_book(book, source={"name": source_name, "sha256": hashlib.sha256(data).hexdigest()},
                             metadata=metadata)


def parse_loaded_book(book, *, source, metadata=None):
    """Parse an in-memory grid from either XLSX or online cell reads; closes it."""
    metrics, issues, blocks = {}, [], {}
    try:
        for key, label, sheet_key, metric in SCENARIOS:
            result = {"label": label, "p99_ms": None, "status": "missing", "evidence": None}
            metrics[key] = result
            block_key = (sheet_key, metric)
            if block_key not in blocks:
                sheets = [sheet for sheet in book.worksheets if compact(sheet.title) == sheet_key]
                try:
                    if len(sheets) != 1:
                        raise InputError(f"需要唯一的 {sheet_key} sheet，实际找到 {len(sheets)} 个。")
                    blocks[block_key] = find_block(sheets[0], metric)
                except InputError as exc:
                    blocks[block_key] = str(exc)
            block = blocks[block_key]
            try:
                if isinstance(block, str):
                    raise InputError(block)
                labels, values, evidence = block
                result["evidence"] = dict(evidence)
                if key.casefold() not in labels:
                    raise InputError(f"header 中缺少场景 {key}。")
                index = labels.index(key.casefold())
                raw = values[index]
                result["evidence"].update({"scenario_index": index, "raw_value": raw})
                if raw in ("-", "", "--", "n/a", "na", "null", "none"):
                    raise InputError("该场景没有 P99 数据。")
                try:
                    number = float(raw)
                except ValueError:
                    raise InputError("P99 不是有效的毫秒数值。") from None
                if not math.isfinite(number) or number < 0:
                    raise InputError("P99 必须是有限、非负的毫秒数值。")
                result.update(p99_ms=number, status="ok")
            except InputError as exc:
                issues.append({"scenario": key, "message": str(exc)})
        count = sum(item["status"] == "ok" for item in metrics.values())
        return {
            "schema_version": 1,
            "status": "ok" if count == 4 else "partial" if count else "failed",
            "unit": "ms",
            "statistic": "P99",
            "source": source,
            "parsed_at": datetime.now(timezone.utc).isoformat(),
            "metadata": metadata or {},
            "metrics": metrics,
            "issues": issues,
        }
    finally:
        book.close()


def comparison_context(metadata, side):
    """Validate the request context needed to compare two result windows."""
    if not isinstance(metadata, dict):
        raise InputError(f"{side} metadata 必须是 JSON 对象。")
    context = {}
    for field in COMPARISON_CONTEXT_FIELDS:
        value = metadata.get(field)
        if not isinstance(value, str) or not value.strip():
            raise InputError(f"{side} metadata 缺少非空 {field}；对比必须明确同一 VIN 的两个时间段。")
        context[field] = value.strip()
    return context


def comparison_number(value):
    """Keep response values numeric, including an exact 0 ms result."""
    return None if value is None else float(value)


def compare_results(baseline, replacement, *, baseline_metadata, replacement_metadata):
    """Compare two independently read AD P99 snapshots without reusing either source."""
    baseline_context = comparison_context(baseline_metadata, "Baseline")
    replacement_context = comparison_context(replacement_metadata, "替换后")
    if baseline_context["vin"] != replacement_context["vin"]:
        raise InputError("Baseline 与替换后 metadata 的 vin 不一致，拒绝生成跨 VIN 对比。")

    comparisons, issues = {}, []
    for key, label, _, _ in SCENARIOS:
        before = baseline.get("metrics", {}).get(key, {})
        after = replacement.get("metrics", {}).get(key, {})
        before_value = comparison_number(before.get("p99_ms"))
        after_value = comparison_number(after.get("p99_ms"))
        row = {
            "label": label,
            "baseline_p99_ms": before_value,
            "replacement_p99_ms": after_value,
            "delta_ms": None,
            "status": "missing",
        }
        if before.get("status") == "ok" and after.get("status") == "ok":
            # Sheet values are decimal text. Decimal avoids response artifacts such as -1.160000000000025.
            row["delta_ms"] = float(Decimal(str(after_value)) - Decimal(str(before_value)))
            row["status"] = "ok"
        else:
            missing = []
            if before.get("status") != "ok":
                missing.append("Baseline")
            if after.get("status") != "ok":
                missing.append("替换后")
            issues.append({"scenario": key, "message": f"{'、'.join(missing)} 缺少可用 P99 数据。"})
        comparisons[key] = row

    count = sum(row["status"] == "ok" for row in comparisons.values())
    return {
        "schema_version": 1,
        "kind": "ad_p99_comparison",
        "status": "ok" if count == len(SCENARIOS) else "partial" if count else "failed",
        "unit": "ms",
        "statistic": "P99",
        "vin": baseline_context["vin"],
        "baseline": {
            "time_range": {"start": baseline_context["start"], "end": baseline_context["end"]},
            "source": baseline.get("source", {}),
        },
        "replacement": {
            "time_range": {"start": replacement_context["start"], "end": replacement_context["end"]},
            "source": replacement.get("source", {}),
        },
        "comparisons": comparisons,
        "table": {
            "columns": ["场景", "指标", "Baseline", "替换小包", "Δ（替换 − Baseline）"],
            "rows": [
                [f"{key}（{row['label']}）", "P99", row["baseline_p99_ms"],
                 row["replacement_p99_ms"], row["delta_ms"]]
                for key, row in comparisons.items()
            ],
        },
        "issues": issues,
    }


def parse_time_token(value):
    """Parse common copied Feishu/log time forms only for ordering two request windows."""
    if re.fullmatch(r"\d{14}", value):
        parsed = datetime.strptime(value, "%Y%m%d%H%M%S")
    elif re.fullmatch(r"\d{12}", value):
        parsed = datetime.strptime(value, "%Y%m%d%H%M")
    else:
        normal = value.replace("/", "-").replace("Z", "+00:00")
        try:
            parsed = datetime.fromisoformat(normal)
        except ValueError:
            raise InputError(f"无法识别时间 {value!r}；请使用 YYYY-MM-DD HH:MM:SS 或 YYYYMMDDHHMMSS。") from None
    if parsed.tzinfo is not None:
        return parsed.astimezone(timezone.utc).replace(tzinfo=None), parsed.isoformat()
    return parsed, parsed.isoformat()


def parse_comparison_request(text):
    """Normalize a copied comparison request before any cards or sheets are read."""
    if not isinstance(text, str) or not text.strip():
        raise InputError("对比请求不能为空。")
    vin_values = [match.group(1) for match in LABELED_VIN_PATTERN.finditer(text)]
    if not vin_values:
        vin_values = [match.group(0) for match in UNLABELED_VIN_PATTERN.finditer(text)]
    vins = {value.strip().upper() for value in vin_values if value.strip()}
    if not vins:
        raise InputError("未识别到 VIN；请写明 VIN 或 VIN 码。")
    if len(vins) != 1:
        raise InputError("对比请求包含多个 VIN，Baseline 与替换后必须是同一 VIN。")
    windows = []
    for match in TIME_RANGE_PATTERN.finditer(text):
        start_at, start = parse_time_token(match.group("start"))
        end_at, end = parse_time_token(match.group("end"))
        if start_at >= end_at:
            raise InputError("每个时间段的起点必须早于终点。")
        windows.append({"start_at": start_at, "end_at": end_at, "start": start, "end": end})
    if len(windows) != 2:
        raise InputError("需要恰好两个时间段；每段使用“起点/终点”或“起点 - 终点”。")
    windows.sort(key=lambda item: (item["start_at"], item["end_at"]))
    baseline, replacement = windows
    if baseline["end_at"] > replacement["start_at"]:
        raise InputError("两个时间段重叠，无法按时间确定刷包前后关系。")
    vin = vins.pop()

    def metadata(window):
        return {"vin": vin, "start": window["start"], "end": window["end"]}

    key_material = json.dumps([vin, baseline["start"], baseline["end"],
                               replacement["start"], replacement["end"]], ensure_ascii=False)
    return {
        "schema_version": 1,
        "kind": "ad_p99_comparison_request",
        "status": "ready",
        "comparison_key": hashlib.sha256(key_material.encode("utf-8")).hexdigest()[:16],
        "vin": vin,
        "baseline_metadata": metadata(baseline),
        "replacement_metadata": metadata(replacement),
        "role_assignment": "较早的时间段为 Baseline，较新的时间段为替换后。",
    }


def inspect_card(document):
    """Inspect supplied card data only. Never open links or execute callback values."""
    buttons = []

    def visit(node):
        if isinstance(node, list):
            for item in node:
                visit(item)
        elif isinstance(node, dict):
            caption = node.get("text", {})
            if isinstance(caption, dict):
                caption = caption.get("content", "")
            if node.get("tag") == "button" and compact(caption) == "查看结果":
                urls = []

                def add(value):
                    if isinstance(value, str) and value.startswith(("https://", "http://")) and value not in urls:
                        urls.append(value)

                add(node.get("url"))
                multi = node.get("multi_url", {})
                if isinstance(multi, dict):
                    for value in multi.values():
                        add(value)
                callback = "value" in node
                for behavior in node.get("behaviors", []) or []:
                    if isinstance(behavior, dict):
                        if behavior.get("type") == "open_url":
                            for key in ("default_url", "url", "pc_url", "ios_url", "android_url"):
                                add(behavior.get(key))
                        elif behavior.get("type") == "callback":
                            callback = True
                buttons.append({"urls": urls, "has_callback": callback})
            for key, value in node.items():
                if key in ("content", "card") and isinstance(value, str):
                    try:
                        visit(json.loads(value))
                    except (ValueError, RecursionError):
                        pass
                elif isinstance(value, (dict, list)):
                    visit(value)

    visit(document)
    return {
        "status": "needs_upstream_confirmation",
        "result_buttons": buttons,
        "message": "这里只识别查看结果按钮；需在公司云电脑确认最终在线表格链接和读取权限，不下载文件。",
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description="公司云电脑专用 AD P99 在线读取；不导出或下载文件。")
    source = parser.add_mutually_exclusive_group()
    source.add_argument("--sheet-url", help="飞书在线表格或指向表格的 wiki 链接；不导出/下载 Excel")
    source.add_argument("--card-json", type=Path, help="仅检查卡片 JSON 的查看结果按钮")
    source.add_argument("--baseline-sheet-url", help="对比用 Baseline 飞书在线表格链接")
    source.add_argument("--comparison-request", type=Path, help="仅解析复制的 VIN 与两个时间段，建立对比上下文")
    parser.add_argument("--replacement-sheet-url", help="对比用替换小包后飞书在线表格链接")
    parser.add_argument("--metadata", type=Path, help="可选 JSON 对象：release、vin、start、end 等原始上下文")
    parser.add_argument("--baseline-metadata", type=Path, help="对比用 Baseline JSON：必须含 vin、start、end")
    parser.add_argument("--replacement-metadata", type=Path, help="对比用替换后 JSON：必须含 vin、start、end")
    parser.add_argument("--output", type=Path, help="输出 JSON 路径；拒绝覆盖已有文件")
    parser.add_argument("--identity", choices=("bot", "user"), help="在线读取必须显式指定获准的飞书身份")
    parser.add_argument("--lark-cli", help="可选：lark-cli 可执行文件路径")
    args = parser.parse_args(argv)
    comparison = bool(args.baseline_sheet_url or args.replacement_sheet_url)
    if not (args.sheet_url or args.card_json or args.comparison_request or comparison):
        parser.error("必须指定 --sheet-url、--card-json、--comparison-request 或一对对比表格链接。")
    if comparison and (args.sheet_url or args.card_json or args.comparison_request):
        parser.error("对比链接不能与 --sheet-url 或 --card-json 同时使用。")
    if comparison and (not args.baseline_sheet_url or not args.replacement_sheet_url):
        parser.error("对比必须同时指定 --baseline-sheet-url 和 --replacement-sheet-url。")
    if comparison and (not args.baseline_metadata or not args.replacement_metadata):
        parser.error("对比必须同时指定 --baseline-metadata 和 --replacement-metadata。")
    if (args.sheet_url or comparison) and args.identity is None:
        parser.error("在线读取必须指定 --identity bot 或 --identity user；不会自动沿用或切换身份。")
    try:
        if args.output and args.output.exists():
            raise InputError("输出文件已存在；请换一个文件名以保留原结果。")
        def load_metadata(path, name):
            metadata = json.loads(read_local(path).decode("utf-8-sig")) if path else {}
            if not isinstance(metadata, dict):
                raise InputError(f"{name} 必须是 JSON 对象。")
            return metadata

        metadata = load_metadata(args.metadata, "metadata")
        if args.sheet_url:
            from lark_source import read_online
            result = read_online(args.sheet_url, identity=args.identity, executable=args.lark_cli, metadata=metadata)
            code = {"ok": 0, "partial": 2, "failed": 1}[result["status"]]
        elif comparison:
            from lark_source import read_online
            baseline_metadata = load_metadata(args.baseline_metadata, "baseline metadata")
            replacement_metadata = load_metadata(args.replacement_metadata, "replacement metadata")
            baseline = read_online(args.baseline_sheet_url, identity=args.identity,
                                   executable=args.lark_cli, metadata=baseline_metadata)
            replacement = read_online(args.replacement_sheet_url, identity=args.identity,
                                      executable=args.lark_cli, metadata=replacement_metadata)
            result = compare_results(baseline, replacement,
                                     baseline_metadata=baseline_metadata,
                                     replacement_metadata=replacement_metadata)
            code = {"ok": 0, "partial": 2, "failed": 1}[result["status"]]
        elif args.comparison_request:
            result = parse_comparison_request(read_local(args.comparison_request).decode("utf-8-sig"))
            code = 0
        elif args.card_json:
            result = inspect_card(json.loads(read_local(args.card_json).decode("utf-8-sig")))
            code = 2
        payload = json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False)
        if args.output:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            with args.output.open("x", encoding="utf-8") as stream:
                stream.write(payload + "\n")
        print(payload)
        return code
    except (InputError, OSError, ValueError, RecursionError) as exc:
        message = str(exc) if isinstance(exc, InputError) else f"输入或输出失败（{type(exc).__name__}）。"
        print(json.dumps({"status": "failed", "error": message}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    # The optional online adapter imports the shared parser and InputError type.
    sys.modules.setdefault("ad_p99", sys.modules[__name__])
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
